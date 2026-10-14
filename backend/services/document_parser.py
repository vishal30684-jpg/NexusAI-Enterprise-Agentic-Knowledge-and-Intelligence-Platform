
from pathlib import Path

import pandas as pd
from docx import Document as WordDocument

from backend.services.pdf_parser import extract_pages_from_pdf


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".csv"}


def extract_pages_from_document(file_path: str) -> list[dict]:
    """
    Extract text from supported document formats.

    Returns a consistent page-based structure for the chunker.
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: {extension or 'unknown'}"
        )

    if extension == ".pdf":
        return extract_pages_from_pdf(str(path))

    if extension in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig").strip()

        return [{"page_number": 1, "text": text}] if text else []

    if extension == ".docx":
        document = WordDocument(str(path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Include text from tables as well.
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells
                )
                if row_text.strip(" |"):
                    paragraphs.append(row_text)

        text = "\n".join(paragraphs)

        return [{"page_number": 1, "text": text}] if text else []

    if extension == ".csv":
        dataframe = pd.read_csv(path, encoding="utf-8-sig")

        # Convert the column names and rows into searchable text.
        text = dataframe.to_string(index=False).strip()

        return [{"page_number": 1, "text": text}] if text else []

    return []
